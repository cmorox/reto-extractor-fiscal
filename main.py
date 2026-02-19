import os
import json
import google.generativeai as genai
from dotenv import load_dotenv
import docx
import openpyxl

# ---------------------------------------------------------
# 1. CONFIGURACIÓN INICIAL
# ---------------------------------------------------------
# Cargar la API Key desde tu archivo .env
load_dotenv()
genai.configure(api_key=os.environ.get("GEMINI_API_KEY"))

# Definir el esquema JSON estricto para la IA
esquema_extraccion = {
    "type": "OBJECT",
    "properties": {
        "nombre_cliente": {"type": "STRING", "description": "El nombre del cliente o empresa."},
        "monto": {"type": "NUMBER", "description": "El monto económico mencionado. Si no hay, devuelve null."},
        "fecha": {"type": "STRING", "description": "La fecha mencionada en formato YYYY-MM-DD."},
        "tipo_solicitud": {
            "type": "STRING",
            "enum": ["Venta", "Queja", "Factura"], 
            "description": "Clasifica la intención del texto en Venta, Queja o Factura."
        }
    },
    "required": ["nombre_cliente", "fecha", "tipo_solicitud"]
}

# Configurar el modelo de Gemini
generation_config = genai.GenerationConfig(
    response_mime_type="application/json",
    response_schema=esquema_extraccion
)
model = genai.GenerativeModel('gemini-2.5-flash', generation_config=generation_config)

# ---------------------------------------------------------
# 2. FUNCIONES DE LECTURA DE ARCHIVOS
# ---------------------------------------------------------
def extraer_texto_de_archivo(ruta_archivo):
    """Detecta el tipo de archivo y extrae su texto en bruto."""
    extension = os.path.splitext(ruta_archivo)[1].lower()
    texto = ""

    try:
        if extension == '.txt':
            with open(ruta_archivo, 'r', encoding='utf-8') as f:
                texto = f.read()
                
        elif extension == '.docx':
            doc = docx.Document(ruta_archivo)
            texto = "\n".join([parrafo.text for parrafo in doc.paragraphs])
            
        elif extension == '.xlsx':
            wb = openpyxl.load_workbook(ruta_archivo, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    # Unir las celdas de la fila que no estén vacías
                    fila = " ".join([str(celda) for celda in row if celda is not None])
                    if fila.strip():
                        texto += fila + "\n"
        else:
            print(f"⚠️ Formato no soportado: {extension}")
            
    except Exception as e:
        print(f"❌ Error al leer {ruta_archivo}: {e}")
        
    return texto.strip()

# ---------------------------------------------------------
# 3. FUNCIÓN DE EXTRACCIÓN CON IA (CON REINTENTOS)
# ---------------------------------------------------------
def procesar_con_ia(texto_sucio, max_reintentos=3):
    """Pasa el texto a Gemini y fuerza la salida en JSON."""
    if not texto_sucio:
        return None
        
    prompt = f"""
    Eres un asistente administrativo experto. Analiza el siguiente texto extraído de un documento.
    Extrae los datos requeridos respetando estrictamente el formato JSON indicado.
    
    Texto a analizar:
    '{texto_sucio}'
    """
    
    for intento in range(max_reintentos):
        try:
            respuesta = model.generate_content(prompt)
            return json.loads(respuesta.text) # Valida que sea JSON correcto
        except json.JSONDecodeError:
            print(f"   [!] Intento {intento + 1} falló (Formato inválido). Reintentando...")
        except Exception as e:
            print(f"   [!] Error de API: {e}")
            
    print("   ❌ Se agotaron los reintentos para este documento.")
    return None

# ---------------------------------------------------------
# 4. FLUJO PRINCIPAL (MAIN)
# ---------------------------------------------------------
if __name__ == "__main__":
    carpeta_entrada = "mis_archivos_sucios"
    archivo_salida = "resultado_final.json"
    
    resultados_globales = []
    
    print("Iniciando El Extractor Fiscal...\n")
    
    # Recorrer todos los archivos en la carpeta
    if os.path.exists(carpeta_entrada):
        archivos = os.listdir(carpeta_entrada)
        
        for nombre_archivo in archivos:
            ruta_completa = os.path.join(carpeta_entrada, nombre_archivo)
            
            # Solo procesar archivos (ignorar subcarpetas si las hubiera)
            if os.path.isfile(ruta_completa):
                print(f"Procesando: {nombre_archivo}")
                
                # Paso A: Leer el texto sucio
                texto_crudo = extraer_texto_de_archivo(ruta_completa)
                
                # Paso B: Extraer con IA
                if texto_crudo:
                    datos_extraidos = procesar_con_ia(texto_crudo)
                    
                    if datos_extraidos:
                        # Le agregamos el nombre del archivo de origen para tener contexto
                        datos_extraidos["archivo_origen"] = nombre_archivo
                        resultados_globales.append(datos_extraidos)
                        print("    Extracción exitosa.")
                else:
                    print("   ⚠️ El archivo estaba vacío o no se pudo leer.")
    else:
        print(f"❌ No se encontró la carpeta: {carpeta_entrada}")

    # Paso C: Guardar todo en resultado_final.json
    if resultados_globales:
        with open(archivo_salida, 'w', encoding='utf-8') as f:
            json.dump(resultados_globales, f, indent=4, ensure_ascii=False)
        print(f"\n¡Proceso terminado! Se han guardado {len(resultados_globales)} registros en '{archivo_salida}'.")
    else:

        print("\n No se extrajeron datos para guardar.")
